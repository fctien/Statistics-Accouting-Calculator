"""重建全部教學產出（一鍵）

用法：python run_all.py

順序
1. python tools/verify_all.py --quick（git 必須乾淨，否則它會拒絕並使本次失敗）
2. python teaching/build_cards.py（Prompt 卡 HTML）
3. python teaching/deck/build_deck.py（四份 pptx；需要 repo 外的 deck_kit，見 README「重建教材」）
4. python teaching/export_pdf.py（PowerPoint COM；PowerPoint 正在執行時回 2，本腳本就停在這裡）
5. python teaching/make_package.py --zip（學生版 zip，缺 PDF 時失敗）

產生 logs/run_all-YYYYMMDD-HHMM.html：每步結果、耗時、輸出末幾行、commit hash，
以及「需要人工目視」的檢查清單（程式無法代勞，log 只列出，不宣稱已完成）
結束碼：0 全部通過；2 在 export_pdf 停下（PowerPoint 正在執行）；1 其他失敗
"""
import html
import os
import sys
import time
from datetime import datetime

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)

from tools.verify_all import CSS, TZ, git, run_cmd  # noqa: E402  共用執行與 log 樣式

TAIL = 15
EXPORT_STEP = "export_pdf（簡報 → PDF）"

# 需要轉成 PDF 後逐頁目視的投影片：(簡報, eyebrow 開頭, 頁標題, 原因)
# 頁碼在執行時由標題找出，簡報增減頁數也不會對錯頁
VISUAL_CHECKS = [
    ("BACalc_AI.pptx", "Step 2", "實際踩過的坑（1／2）", "Task 30 估算最密的頁之一"),
    ("BACalc_AI.pptx", "Step 4", "Prompt 原文（1／2）", "Task 30 估算最密的頁之一"),
    ("BACalc_AI.pptx", "Step 4", "實際踩過的坑（1／2）", "Task 30 估算最密的頁之一"),
    ("BACalc_AI.pptx", "Step 8", "Prompt 原文（1／2）", "Task 30 估算最密的頁之一"),
    ("BACalc_AI.pptx", "Step 8", "實際踩過的坑（1／2）", "Task 30 估算最密的頁之一"),
    ("BACalc_AI.pptx", "Step 0", "一句話下法與四段式對照", "Task 31 改成實測文字"),
    ("BACalc_AI.pptx", "Step 1", "一句話下法與四段式對照", "Task 31 改成實測文字"),
    ("BACalc_AI.pptx", "Step 2", "一句話下法與四段式對照", "Task 31 改成實測文字"),
    ("BACalc_AI.pptx", "Step 0", "實際踩過的坑（2／2）", "Task 31 新增演練的坑"),
    ("BACalc_AI.pptx", "Step 1", "實際踩過的坑（2／2）", "Task 31 新增演練的坑"),
    ("BACalc_AI.pptx", "Step 2", "實際踩過的坑（2／2）", "Task 31 新增演練的坑"),
    ("BACalc_EMBA.pptx", "", "專案層級範本（1／4）", "Task 30 估算最密的頁之一"),
    ("BACalc_EMBA.pptx", "Step 9", "怎麼對答案", "Task 30 估算最密的頁之一"),
]


def steps():
    py = sys.executable
    return [
        ("verify_all --quick", [py, "tools/verify_all.py", "--quick"]),
        ("build_cards（Prompt 卡）", [py, "teaching/build_cards.py"]),
        ("build_deck（簡報 pptx）", [py, "teaching/deck/build_deck.py"]),
        (EXPORT_STEP, [py, "teaching/export_pdf.py"]),
        ("make_package --zip（學生版）", [py, "teaching/make_package.py", "--zip"]),
    ]


def slide_texts(path):
    """回傳每頁前三個文字框的文字（eyebrow、標題、副標）；讀不到回 None"""
    try:
        from pptx import Presentation
        prs = Presentation(path)
    except Exception:
        return None
    pages = []
    for s in prs.slides:
        pages.append([sh.text_frame.text for sh in s.shapes if sh.has_text_frame][:3])
    return pages


def find_visual_pages(out_dir):
    """VISUAL_CHECKS → [(簡報, 頁碼或「找不到」, 頁標題, 原因)]"""
    cache = {}
    rows = []
    for deck, eyebrow, title, why in VISUAL_CHECKS:
        if deck not in cache:
            cache[deck] = slide_texts(os.path.join(out_dir, deck))
        pages = cache[deck] or []
        hits = [i for i, t in enumerate(pages, 1)
                if title in t[1:3] and (not eyebrow or (t and t[0].startswith(eyebrow)))]
        rows.append((deck, ", ".join(map(str, hits)) if hits else "找不到", "%s %s" % (eyebrow, title), why))
    return rows


def render(meta, results, not_run, status, checklist):
    e = html.escape
    rows = []
    for name, cmd, code, secs, out in results:
        cls, label = ("ok", "PASS") if code == 0 else (
            ("ng", "停止（exit 2）") if code == 2 and name == EXPORT_STEP else ("ng", "FAIL（exit %d）" % code))
        tail = "\n".join(out.splitlines()[-TAIL:])
        rows.append("<tr><td>%s</td><td><code>%s</code></td><td class=%s>%s</td><td>%.1f 秒</td>"
                    "<td><pre>%s</pre></td></tr>" % (e(name), e(cmd), cls, e(label), secs, e(tail)))
    for name in not_run:
        rows.append("<tr><td>%s</td><td>—</td><td>未執行</td><td>—</td><td>前一步停止</td></tr>" % e(name))
    meta_rows = "".join("<tr><th>%s</th><td>%s</td></tr>" % (e(k), e(v)) for k, v in meta)
    checks = "".join(
        "<tr><td><input type=checkbox disabled> 未完成</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>"
        % (e(d), e(p), e(t), e(w)) for d, p, t, w in checklist)
    return (
        '<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        "<title>BACalc 教材重建</title><style>%s</style></head><body>"
        "<h1>BACalc 教材重建（run_all）</h1><p class=%s>%s</p><table>%s</table>"
        "<h2>各步驟</h2><table><tr><th>步驟</th><th>指令</th><th>結果</th><th>耗時</th>"
        "<th>輸出（末 %d 行）</th></tr>%s</table>"
        "<h2>人工目視檢查清單（程式無法代勞，本 log 不宣稱已完成）</h2>"
        "<p>PDF 匯出後，逐頁打開下列投影片確認文字沒有超出框、沒有被裁切、行距正常；"
        "頁碼依本次 pptx 的標題自動找出</p>"
        "<table><tr><th>狀態</th><th>簡報</th><th>頁碼</th><th>頁面</th><th>原因</th></tr>%s</table>"
        "</body></html>"
        % (CSS, "ok" if status == 0 else "ng", e(meta[-1][1]), meta_rows, TAIL, "".join(rows), checks))


def main(log_dir=None, runner=run_cmd, out_dir=None):
    """Controller；log_dir、runner、out_dir 可注入以便測試"""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")
    started = datetime.now(TZ)
    head = git("rev-parse", "HEAD")
    dirty_before = git("status", "--porcelain")
    results, not_run, status = [], [], 0
    plan = steps()
    for i, (name, cmd) in enumerate(plan):
        print("==> %s" % name, flush=True)
        t0 = time.time()
        code, out = runner(cmd)
        secs = time.time() - t0
        shown = " ".join(["python"] + cmd[1:])
        results.append((name, shown, code, secs, out))
        print("    exit %d（%.1f 秒）" % (code, secs), flush=True)
        if code != 0:
            print("\n".join(out.splitlines()[-TAIL:]))
            status = 2 if (code == 2 and name == EXPORT_STEP) else 1
            not_run = [n for n, _ in plan[i + 1:]]
            break
    summary = {0: "總結：PASS（%d 步全過）" % len(results),
               1: "總結：FAIL（%s 失敗）" % results[-1][0],
               2: "總結：停在 export_pdf（PowerPoint 正在執行，未匯出 PDF，學生版 zip 未產生）"}[status]
    meta = [
        ("時間", started.strftime("%Y-%m-%d %H:%M:%S +08:00")),
        ("commit", "%s（短碼 %s）" % (head, head[:7])),
        ("分支", git("rev-parse", "--abbrev-ref", "HEAD")),
        ("執行前 git 狀態", "乾淨" if not dirty_before else "不乾淨（verify_all 會拒絕）"),
        ("Python", sys.version.split()[0]),
        ("結果", summary),
    ]
    out_dir = out_dir or os.path.join(ROOT, "teaching", "out")
    page = render(meta, results, not_run, status, find_visual_pages(out_dir))
    log_dir = log_dir or os.path.join(ROOT, "logs")
    os.makedirs(log_dir, exist_ok=True)
    path = os.path.join(log_dir, "run_all-%s.html" % started.strftime("%Y%m%d-%H%M"))
    with open(path, "w", encoding="utf-8") as f:
        f.write(page)
    print("log：%s" % path)
    print(summary)
    return status


if __name__ == "__main__":
    sys.exit(main())
